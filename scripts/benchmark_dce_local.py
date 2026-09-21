"""Bounded, synthetic local benchmark for deterministic DCE extraction."""

from __future__ import annotations

import argparse
import json
import statistics
import time
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.modules.dce.application.extraction import _project_document


def _percentile(values: list[float], percentile: float) -> float:
    ordered = sorted(values)
    return ordered[min(len(ordered) - 1, round((len(ordered) - 1) * percentile))]


def _text_bytes() -> bytes:
    return "\n".join(
        f"Article {index}: exigence technique, délai et pièce justificative." for index in range(20_000)
    ).encode()


def _docx_bytes() -> bytes:
    document = Document()
    for index in range(2_000):
        document.add_paragraph(f"Article DOCX {index}: exigence et réserve à vérifier.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "BPU"
    for sheet_index in range(3):
        sheet = first if sheet_index == 0 else workbook.create_sheet(f"Lot-{sheet_index}")
        for row in range(1, 201):
            for column in range(1, 21):
                sheet.cell(row=row, column=column, value=f"Lot {sheet_index} cellule {row}-{column}")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--iterations", type=int, default=3)
    args = parser.parse_args()
    if args.iterations < 1:
        raise SystemExit("--iterations must be positive")

    corpus = {
        "text/plain": _text_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _docx_bytes(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _xlsx_bytes(),
    }
    results: dict[str, object] = {
        "benchmark": "SMART_AO_LOCAL_DCE_EXTRACTION_V0",
        "scope": "synthetic bounded corpus; deterministic native extraction only",
        "iterations": args.iterations,
        "corpus": {},
    }
    for media_type, source_bytes in corpus.items():
        durations: list[float] = []
        projection = None
        for _ in range(args.iterations):
            started = time.perf_counter()
            projection = _project_document(media_type=media_type, source_bytes=source_bytes)
            durations.append((time.perf_counter() - started) * 1000)
        assert projection is not None
        results["corpus"][media_type] = {
            "bytes": len(source_bytes),
            "status": projection.status,
            "fragments": len(projection.fragments),
            "characters": sum(len(fragment.text) for fragment in projection.fragments),
            "latency_ms": {
                "min": round(min(durations), 3),
                "median": round(statistics.median(durations), 3),
                "p95": round(_percentile(durations, 0.95), 3),
                "max": round(max(durations), 3),
            },
        }
        if projection.status != "COMPLETED":
            raise RuntimeError(f"unexpected extraction status for {media_type}: {projection.status}")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(results, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(results, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
